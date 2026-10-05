from contextvars import ContextVar
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from models.db import Report
from services.exports import slide_layout as L
from services.exports.helpers import logo_path
from services.exports.slides import (
    BODY_TOP,
    CONTENT_WIDTH,
    HEADING_ROW_HEIGHT,
    MARGIN_X,
    SLIDE_HEIGHT,
    SLIDE_WIDTH,
    TEXT_BOX_PADDING,
    Slide,
    SlideDeck,
    build_deck,
    text_line_height,
)
from services.exports.themes import (
    BORDER_HEX,
    DEFAULT_THEME,
    MUTED_HEX,
    SURFACE_HEX,
    SURFACE_STRONG_HEX,
    TEXT_HEX,
    WHITE_HEX,
    ReportTheme,
)

MUTED = MUTED_HEX
TEXT = TEXT_HEX
WHITE = WHITE_HEX
BORDER = BORDER_HEX
SURFACE = SURFACE_HEX
SURFACE_STRONG = SURFACE_STRONG_HEX
TAG_H2 = MUTED_HEX
EMU_PER_INCH = 914_400
NO_STYLE_TABLE_ID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
CARD_RADIUS = 0.08
BADGE_RADIUS = 0.1
TAG_RADIUS = 0.05

_theme: ContextVar[ReportTheme] = ContextVar("pptx_theme", default=DEFAULT_THEME)


def render_report_pptx(report: Report) -> bytes:
    deck = build_deck(report, "png")
    token = _theme.set(deck.theme)
    try:
        return _render(deck)
    finally:
        _theme.reset(token)


def _render(deck: SlideDeck) -> bytes:
    presentation = Presentation()
    presentation.slide_width = _in(SLIDE_WIDTH)
    presentation.slide_height = _in(SLIDE_HEIGHT)
    presentation.core_properties.title = f"{deck.site_name} · {deck.report_title}"
    blank_layout = presentation.slide_layouts[6]

    for item in deck.slides:
        slide = presentation.slides.add_slide(blank_layout)
        builder = _BUILDERS[item.kind]
        if item.kind not in ("title", "closing"):
            _chrome(slide, deck, item)
        builder(slide, deck, item)

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


# ---------- Slides ----------
def _title_slide(slide, deck: SlideDeck, _: Slide) -> None:
    _rect(slide, L.TITLE_STRIP, _primary())
    _logo(slide, L.TITLE_LOGO)
    _text(slide, L.TITLE_EYEBROW, deck.report_title.upper(), 13, bold=True, color=_accent(),
          letter_spacing=2)
    _text(slide, L.TITLE_SITE, deck.site_name, 44, bold=True, color=_primary())
    _text(slide, L.TITLE_URL, deck.source_url, 16, color=MUTED)
    _rect(slide, L.TITLE_BAR, _accent())
    plural = "s" if deck.page_count != 1 else ""
    _text(slide, L.TITLE_META,
          f"Generated {deck.generated_at}   ·   {deck.page_count} page{plural} analysed", 13,
          color=MUTED)
    _text(slide, L.TITLE_BRAND, deck.brand_name, 12, bold=True, color=_primary())


def _overview_slide(slide, deck: SlideDeck, item: Slide) -> None:
    _slide_title(slide, item.title)
    metrics = (
        ("Pages analysed", f"{deck.page_count}"),
        ("Total words", f"{deck.totals.words:,}"),
        ("Total links", f"{deck.totals.links:,}"),
        ("Total images", f"{deck.totals.images:,}"),
    )
    for index, (label, value) in enumerate(metrics):
        x = MARGIN_X + index * (L.METRIC_WIDTH + L.METRIC_GAP)
        _card(slide, (x, L.METRIC_TOP, L.METRIC_WIDTH, L.METRIC_HEIGHT))
        _text(slide, (x + 0.25, L.METRIC_TOP + 0.27, L.METRIC_WIDTH - 0.4, 0.3), label.upper(),
              11, color=MUTED, letter_spacing=0.8)
        _text(slide, (x + 0.25, L.METRIC_TOP + 0.6, L.METRIC_WIDTH - 0.4, 0.55), value, 30,
              bold=True, color=_primary())

    x, y, width, height = L.DETAILS
    _rect(slide, L.DETAILS, SURFACE, radius=CARD_RADIUS)
    details = (
        ("Website read on", deck.created_at),
        ("Report generated", deck.generated_at),
        ("Shareable link", deck.share_url),
    )
    for index, (label, value) in enumerate(details):
        row_y = y + L.DETAILS_ROW_TOP + index * L.DETAILS_ROW_STEP
        _text(slide, (x + L.DETAILS_LABEL_X, row_y, 2.7, 0.32), label, 14, bold=True, color=MUTED)
        _text(slide, (x + L.DETAILS_VALUE_X, row_y, width - L.DETAILS_VALUE_X - 0.3, 0.32), value,
              14, color=_primary() if index == 2 else TEXT)


def _summary_slide(slide, deck: SlideDeck, item: Slide) -> None:
    _slide_title(slide, item.title)
    rows = item.data["rows"]
    show_total = item.data["show_total"]
    row_count = 1 + len(rows) + (1 if show_total else 0)
    height = (L.TABLE_HEADER_HEIGHT + len(rows) * L.TABLE_ROW_HEIGHT
              + (L.TABLE_TOTAL_HEIGHT if show_total else 0))
    frame = slide.shapes.add_table(row_count, len(L.TABLE_COLUMNS), _in(MARGIN_X),
                                   _in(L.TABLE_TOP), _in(CONTENT_WIDTH), _in(height))
    frame._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId")).text = (
        NO_STYLE_TABLE_ID
    )
    table = frame.table
    for column, width in zip(table.columns, L.TABLE_COLUMNS):
        column.width = _in(width)

    headers = ("#", "Page", "Description", "Words", "Links", "Images")
    table.rows[0].height = _in(L.TABLE_HEADER_HEIGHT)
    for index, text in enumerate(headers):
        cell = table.cell(0, index)
        _cell(cell, fill=_primary())
        _cell_text(cell, text, 11, bold=True, color=WHITE, align=_column_align(index))

    for row_index, row in enumerate(rows, start=1):
        table.rows[row_index].height = _in(L.TABLE_ROW_HEIGHT)
        fill = SURFACE if row_index % 2 == 0 else WHITE
        values = (str(row["position"]), None, row["description"], f"{row['words']:,}",
                  f"{row['links']:,}", f"{row['images']:,}")
        for index, value in enumerate(values):
            cell = table.cell(row_index, index)
            _cell(cell, fill=fill, bottom=BORDER)
            if value is not None:
                _cell_text(cell, value, 10 if index == 2 else 10.5, align=_column_align(index))
        page_cell = table.cell(row_index, 1)
        _cell_text(page_cell, row["title"], 10.5, bold=True, color=_primary())
        _add_paragraph(page_cell.text_frame, row["url"], 9, color=MUTED)

    if show_total:
        last = row_count - 1
        table.rows[last].height = _in(L.TABLE_TOTAL_HEIGHT)
        totals = ("Total", "", "", f"{deck.totals.words:,}", f"{deck.totals.links:,}",
                  f"{deck.totals.images:,}")
        for index, value in enumerate(totals):
            cell = table.cell(last, index)
            _cell(cell, fill=SURFACE_STRONG, top=_primary())
            if value:
                align = PP_ALIGN.LEFT if index == 0 else _column_align(index)
                _cell_text(cell, value, 10.5, bold=True, color=_primary(), align=align)
        table.cell(last, 0).merge(table.cell(last, 2))


def _chart_slide(slide, _: SlideDeck, item: Slide) -> None:
    chart = item.data["chart"]
    _slide_title(slide, item.title)
    _text(slide, L.CHART_CAPTION, chart.caption, 13, color=MUTED)
    _picture(slide, chart.image, L.CHART_IMAGE)


def _page_slide(slide, _: SlideDeck, item: Slide) -> None:
    data = item.data
    _page_heading(slide, data)
    for (x, width), (label, value) in zip(L.STAT_COLUMNS, data["stats"]):
        _rect(slide, (x, L.STAT_TOP, width, L.STAT_HEIGHT), SURFACE_STRONG, radius=CARD_RADIUS)
        _text(slide, (x + 0.2, L.STAT_TOP + 0.16, width - 0.3, 0.26), label.upper(), 10,
              bold=True, color=MUTED, letter_spacing=0.8)
        _text(slide, (x + 0.2, L.STAT_TOP + 0.45, width - 0.3, 0.38), value, 18,
              bold=True, color=_primary(), wrap=False)

    if data["chart"]:
        _text(slide, L.BREAKDOWN_TITLE, data["chart"].title, 13, bold=True, color=_primary())
        _picture(slide, data["chart"].image, L.BREAKDOWN_IMAGE)
    if data["meta"] or data["meta_note"]:
        _text(slide, L.META_TITLE, "Meta description", 13, bold=True, color=_primary())
        _text(slide, L.META_TEXT, data["meta"] or data["meta_note"], 11,
              color=TEXT if data["meta"] else MUTED, italic=not data["meta"],
              line_height=_theme.get().pt(16))


def _page_content_slide(slide, _: SlideDeck, item: Slide) -> None:
    data = item.data
    _page_heading(slide, data)
    for section in data["sections"]:
        top = BODY_TOP + section["top"]
        _text(slide, (MARGIN_X, top, CONTENT_WIDTH, 0.32), section["title"], 13, bold=True,
              color=_primary())
        body_top = BODY_TOP + section["body_top"]
        if section["kind"] == "headings":
            _heading_rows(slide, section, body_top)
        else:
            _text_box(slide, section, body_top)


def _closing_slide(slide, deck: SlideDeck, _: Slide) -> None:
    _logo(slide, L.CLOSING_LOGO)
    _text(slide, L.CLOSING_TITLE, "Thank you", 40, bold=True, color=_primary(), align=PP_ALIGN.CENTER)
    _rect(slide, L.CLOSING_BAR, _accent())
    _text(slide, L.CLOSING_LEAD, "View the full report online", 14, color=MUTED,
          align=PP_ALIGN.CENTER)
    _text(slide, L.CLOSING_LINK, deck.share_url, 15, bold=True, color=_primary(),
          align=PP_ALIGN.CENTER)
    _text(slide, L.CLOSING_NOTE, f"{deck.brand_name}   ·   Generated {deck.generated_at}", 11,
          color=MUTED, align=PP_ALIGN.CENTER)


_BUILDERS = {
    "title": _title_slide,
    "overview": _overview_slide,
    "summary": _summary_slide,
    "chart": _chart_slide,
    "page": _page_slide,
    "page_content": _page_content_slide,
    "closing": _closing_slide,
}


# ---------- Shared pieces ----------
def _chrome(slide, deck: SlideDeck, item: Slide) -> None:
    _logo(slide, L.HEADER_LOGO)
    _text(slide, L.HEADER_BRAND, deck.brand_name, 12, bold=True, color=_primary(),
          anchor=MSO_ANCHOR.MIDDLE)
    _text(slide, L.HEADER_SITE, deck.site_name, 11, bold=True, color=_primary(),
          align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    _line(slide, MARGIN_X, L.HEADER_LINE_Y, L.RIGHT_EDGE, BORDER)
    _line(slide, MARGIN_X, L.FOOTER_LINE_Y, L.RIGHT_EDGE, BORDER)
    _text(slide, L.FOOTER_TEXT, f"{deck.site_name}   ·   {deck.report_title}", 9, color=MUTED,
          anchor=MSO_ANCHOR.MIDDLE)
    _text(slide, L.FOOTER_NUMBER, f"{item.number} / {len(deck.slides)}", 9, color=MUTED,
          align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def _slide_title(slide, title: str) -> None:
    _text(slide, L.SLIDE_TITLE, title, 24, bold=True, color=_primary(), anchor=MSO_ANCHOR.MIDDLE)
    _rect(slide, L.SLIDE_TITLE_BAR, _accent())


def _page_heading(slide, data: dict) -> None:
    badge = _rect(slide, L.PAGE_BADGE, _primary(), radius=BADGE_RADIUS)
    _shape_text(badge, f"{data['position']:02d}", 18, bold=True, color=WHITE)
    _text(slide, L.PAGE_TITLE, data["title"], 20, bold=True, color=_primary(),
          anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    _text(slide, L.PAGE_URL, data["url"], 11, color=MUTED, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
    _rect(slide, L.PAGE_RULE, _accent())


def _card(slide, box) -> None:
    x, y, width, _ = box
    _rect(slide, box, SURFACE, line=BORDER, radius=CARD_RADIUS)
    _rect(slide, (x + CARD_RADIUS, y, width - 2 * CARD_RADIUS, 0.05), _accent())


def _heading_rows(slide, section: dict, body_top: float) -> None:
    if section.get("note"):
        _text(slide, (MARGIN_X, body_top, CONTENT_WIDTH, HEADING_ROW_HEIGHT), section["note"],
              11.5, color=MUTED, italic=True, anchor=MSO_ANCHOR.MIDDLE)
    tag_width, tag_height = L.TAG_SIZE
    for index, heading in enumerate(section["items"]):
        row_y = body_top + index * HEADING_ROW_HEIGHT
        indent = L.H2_INDENT if heading["level"] == 2 else 0.0
        tag = _rect(slide, (MARGIN_X + indent, row_y + (HEADING_ROW_HEIGHT - tag_height) / 2,
                            tag_width, tag_height),
                    _accent() if heading["level"] == 1 else TAG_H2, radius=TAG_RADIUS)
        _shape_text(tag, f"H{heading['level']}", 8, bold=True, color=WHITE)
        text_x = MARGIN_X + indent + L.HEADING_TEXT_OFFSET
        _text(slide, (text_x, row_y, L.RIGHT_EDGE - text_x, HEADING_ROW_HEIGHT), heading["text"],
              11.5, anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        _line(slide, MARGIN_X, row_y + HEADING_ROW_HEIGHT, L.RIGHT_EDGE, BORDER, dash=True)


def _text_box(slide, section: dict, body_top: float) -> None:
    height = section["body_height"]
    _rect(slide, (MARGIN_X, body_top, CONTENT_WIDTH, height), SURFACE)
    _rect(slide, (MARGIN_X, body_top, L.TEXT_BAR_WIDTH, height), _accent())
    _text(slide, (MARGIN_X + L.TEXT_INSET_X, body_top + TEXT_BOX_PADDING,
                  CONTENT_WIDTH - 2 * L.TEXT_INSET_X, height - 2 * TEXT_BOX_PADDING),
          section["text"], 12, line_height=text_line_height(_theme.get()) * 72)


# ---------- Low-level helpers ----------
def _in(value: float) -> Emu:
    return Emu(round(value * EMU_PER_INCH))


def _rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def _primary() -> str:
    return _theme.get().primary


def _accent() -> str:
    return _theme.get().accent


def _font(font, size: float, bold: bool, color: str, italic: bool = False) -> None:
    theme = _theme.get()
    font.name = theme.font_name
    font.size = Pt(theme.pt(size))
    font.bold = bold
    font.italic = italic
    font.color.rgb = _rgb(color)


def _text(slide, box, text: str, size: float, *, bold: bool = False, italic: bool = False,
          color: str = TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap: bool = True,
          line_height: float | None = None, letter_spacing: float | None = None):
    x, y, width, height = box
    shape = slide.shapes.add_textbox(_in(x), _in(y), _in(width), _in(height))
    frame = shape.text_frame
    frame.word_wrap = wrap
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    if line_height:
        paragraph.line_spacing = Pt(line_height)
    run = paragraph.add_run()
    run.text = text
    _font(run.font, size, bold, color, italic)
    if letter_spacing:
        run.font._element.set("spc", str(round(letter_spacing * 100)))
    return shape


def _rect(slide, box, fill: str | None, *, line: str | None = None, radius: float | None = None):
    x, y, width, height = box
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, _in(x), _in(y), _in(width), _in(height))
    if radius:
        shape.adjustments[0] = min(radius / min(width, height), 0.5)
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _shape_text(shape, text: str, size: float, *, bold: bool, color: str) -> None:
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.word_wrap = False
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    _font(run.font, size, bold, color)


def _line(slide, x1: float, y: float, x2: float, color: str, dash: bool = False) -> None:
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _in(x1), _in(y), _in(x2), _in(y))
    # The default connector style references a themed line and shadow; drop it so only ours applies.
    style = connector._element.find(qn("p:style"))
    if style is not None:
        connector._element.remove(style)
    connector.line.color.rgb = _rgb(color)
    connector.line.width = Pt(0.75)
    if dash:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def _picture(slide, image: bytes, box) -> None:
    x, y, width, height = box
    slide.shapes.add_picture(BytesIO(image), _in(x), _in(y), _in(width), _in(height))


def _logo(slide, box) -> None:
    path = logo_path()
    if path:
        x, y, width, height = box
        slide.shapes.add_picture(str(path), _in(x), _in(y), _in(width), _in(height))


def _cell(cell, *, fill: str, top: str | None = None, bottom: str | None = None) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(fill)
    cell.margin_left = cell.margin_right = _in(0.1)
    cell.margin_top = cell.margin_bottom = _in(0.05)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tc_pr = cell._tc.get_or_add_tcPr()
    # Borders must come first in a:tcPr, in schema order: left, right, top, bottom.
    borders = (("a:lnL", None, 0), ("a:lnR", None, 0), ("a:lnT", top, 1.5), ("a:lnB", bottom, 0.75))
    for index, (tag, color, width) in enumerate(borders):
        line = tc_pr.makeelement(qn(tag), {"w": str(Pt(width) if color else 0)})
        if color:
            solid = line.makeelement(qn("a:solidFill"), {})
            solid.append(solid.makeelement(qn("a:srgbClr"), {"val": color}))
            line.append(solid)
        else:
            line.append(line.makeelement(qn("a:noFill"), {}))
        tc_pr.insert(index, line)


def _cell_text(cell, text: str, size: float, *, bold: bool = False, color: str = TEXT,
               align=PP_ALIGN.LEFT) -> None:
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    _font(run.font, size, bold, color)


def _add_paragraph(frame, text: str, size: float, *, color: str) -> None:
    paragraph = frame.add_paragraph()
    run = paragraph.add_run()
    run.text = text
    _font(run.font, size, False, color)


def _column_align(index: int):
    if index == 0:
        return PP_ALIGN.CENTER
    return PP_ALIGN.RIGHT if index >= 3 else PP_ALIGN.LEFT
