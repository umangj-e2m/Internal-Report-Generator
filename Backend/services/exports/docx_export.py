from contextvars import ContextVar
from io import BytesIO

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsmap, qn
from docx.section import Section
from docx.shared import Emu, Mm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from config import get_settings
from models.db import Report, ReportPage
from services.exports.charts import ReportChart, build_charts, build_page_chart
from services.exports.helpers import compute_totals, generated_at_label, logo_path, page_description
from services.exports.themes import (
    BORDER_HEX,
    DEFAULT_THEME,
    MUTED_HEX,
    SURFACE_HEX,
    SURFACE_STRONG_HEX,
    TEXT_HEX,
    WHITE_HEX,
    ReportTheme,
    theme_for,
)
from services.reports.helpers import share_url
from services.shared.constants import REPORT_TITLE
from utils.dates import format_display

MUTED = RGBColor.from_string(MUTED_HEX)
WHITE = RGBColor.from_string(WHITE_HEX)

PAGE_WIDTH = Mm(210)
PAGE_HEIGHT = Mm(297)
SIDE_MARGIN = Mm(16)
CONTENT_WIDTH = Emu(PAGE_WIDTH - 2 * SIDE_MARGIN)
SUMMARY_COLUMN_WIDTHS = (Mm(9), Mm(54), Mm(70), Mm(15), Mm(13), Mm(17))
HEADER_COLUMN_WIDTHS = (Mm(12), Mm(80), Emu(CONTENT_WIDTH - Mm(92)))
METRIC_GAP = Mm(3.5)
METRIC_CARD_HEIGHT = Mm(17.5)
DETAILS_BOX_HEIGHT = Mm(27)
DETAILS_BOX_RADIUS = 7_000
DETAILS_VALUE_OFFSET = Mm(37.5)
STATS_COLUMN_SHARES = (0.22, 0.22, 0.22, 0.34)
TAG_SIZE = (Mm(6.9), Mm(4.2))
TAG_RADIUS = 19_000
TAG_BASELINE_SHIFT_HALF_PT = -5
SUMMARY_PADDING = (150, 150, 180, 180)  # top, bottom, left, right in dxa (7.5pt / 9pt)
METRIC_CARD_RADIUS = 10_000  # DrawingML roundRect "adj": 1/100000 of the shorter side
METRIC_ACCENT_PCT = 5_000  # accent top band as 1/100000 of the card height (~3px)
PAGE_BADGE_SIZE = (Mm(11), Mm(8.5))
PAGE_BADGE_RADIUS = 18_000
PAGE_BADGE_WIDTH = Mm(14)
PAGE_TITLE_TOP_SPACE = Pt(12)
HEADING_H2_INDENT = Mm(4.8)
WPS_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"

_theme: ContextVar[ReportTheme] = ContextVar("docx_theme", default=DEFAULT_THEME)

# OOXML requires child elements in schema order; these are the siblings that must come after.
_PPR_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
_TCPR_AFTER_BORDERS = (
    "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
    "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange",
)
_TCPR_AFTER_SHD = _TCPR_AFTER_BORDERS[1:]
_RPR_AFTER_SHD = (
    "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
    "w:specVanish", "w:oMath",
)
_RPR_AFTER_POSITION = (
    "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", *_RPR_AFTER_SHD,
)
_RPR_AFTER_SPACING = ("w:w", "w:kern", "w:position", *_RPR_AFTER_POSITION)
_TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                        "w:tblDescription", "w:tblPrChange")
_TBLPR_AFTER_CELL_MARGINS = ("w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def render_report_docx(report: Report) -> bytes:
    token = _theme.set(theme_for(report))
    try:
        return _render(report)
    finally:
        _theme.reset(token)


def _render(report: Report) -> bytes:
    doc = Document()
    _setup_styles(doc)

    section = doc.sections[0]
    _setup_page(section)
    _build_header(section, report.site_name)
    _build_footer(section, report.site_name)

    _build_cover(doc, report)
    _build_summary(doc, report)
    _build_charts(doc, report)
    for page in report.pages:
        _build_page_section(doc, page)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------- Layout ----------
def _setup_page(section: Section) -> None:
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = Mm(26)
    section.bottom_margin = Mm(20)
    section.left_margin = SIDE_MARGIN
    section.right_margin = SIDE_MARGIN
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)


def _setup_styles(doc: DocxDocument) -> None:
    theme = _theme.get()
    normal = doc.styles["Normal"]
    _set_style_font(normal)
    normal.font.size = Pt(theme.pt(10.5))
    normal.font.color.rgb = RGBColor.from_string(TEXT_HEX)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    heading1 = doc.styles["Heading 1"]
    _set_style_font(heading1)
    heading1.font.size = Pt(theme.pt(15))
    heading1.font.bold = True
    heading1.font.color.rgb = _primary()
    heading1.paragraph_format.space_before = Pt(0)
    heading1.paragraph_format.space_after = Pt(6)

    heading2 = doc.styles["Heading 2"]
    _set_style_font(heading2)
    heading2.font.size = Pt(theme.pt(11.5))
    heading2.font.bold = True
    heading2.font.color.rgb = _primary()
    heading2.paragraph_format.space_before = Pt(12)
    heading2.paragraph_format.space_after = Pt(4)


def _set_style_font(style) -> None:
    font_name = _theme.get().font_name
    style.font.name = font_name
    r_fonts = style.element.rPr.rFonts
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), font_name)
    # Built-in styles point at theme fonts, which take precedence over explicit names.
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        r_fonts.attrib.pop(qn(attribute), None)


def _build_header(section: Section, site_name: str) -> None:
    header = section.header
    table = header.add_table(rows=1, cols=3, width=CONTENT_WIDTH)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table, None, margins=(0, 110, 0, 0))
    _set_column_widths(table, HEADER_COLUMN_WIDTHS)

    # A header must end with a paragraph, so the table goes before the default one, which is collapsed.
    trailing = header.paragraphs[0]
    trailing._p.addprevious(table._tbl)
    _collapse_paragraph(trailing)

    logo_cell, brand_cell, title_cell = table.rows[0].cells
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_border(cell, "bottom", BORDER_HEX, size=6)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        cell.paragraphs[0].paragraph_format.line_spacing = 1.0

    logo = logo_path()
    if logo:
        logo_cell.paragraphs[0].add_run().add_picture(str(logo), width=Mm(9), height=Mm(9))
    _run(brand_cell.paragraphs[0], get_settings().report_brand_name, size=10, bold=True,
         color=_primary())

    title_line = title_cell.paragraphs[0]
    title_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(title_line, REPORT_TITLE, size=8, color=MUTED)
    site_line = title_cell.add_paragraph()
    site_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    site_line.paragraph_format.space_after = Pt(0)
    site_line.paragraph_format.line_spacing = 1.0
    _run(site_line, site_name, size=8, bold=True, color=_primary())


def _build_footer(section: Section, site_name: str) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.style = "Normal"
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    _set_paragraph_border(paragraph, "top")

    _run(paragraph, f"{site_name} · Generated {generated_at_label()}\tPage ", size=8, color=MUTED)
    _add_field(paragraph, "PAGE")
    _run(paragraph, " of ", size=8, color=MUTED)
    _add_field(paragraph, "NUMPAGES")


# ---------- Content ----------
def _build_cover(doc: DocxDocument, report: Report) -> None:
    settings = get_settings()
    totals = compute_totals(report)

    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_after = Pt(2)
    _letter_spacing(_run(eyebrow, REPORT_TITLE.upper(), size=8.5, bold=True, color=_accent()), 1.2)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    _run(title, report.site_name, size=24, bold=True, color=_primary())

    url = doc.add_paragraph()
    url.paragraph_format.space_after = Pt(12)
    _run(url, report.source_url, size=10, color=MUTED)

    metrics = [
        ("Pages analysed", f"{report.page_count}"),
        ("Total words", f"{totals.words:,}"),
        ("Total links", f"{totals.links:,}"),
        ("Total images", f"{totals.images:,}"),
    ]
    # Separate cards: card columns alternate with narrow, unstyled gap columns.
    cols = len(metrics) * 2 - 1
    table = _new_table(doc, rows=1, cols=cols, borders=False)
    _set_table_borders(table, None, margins=(0, 0, 0, 0))
    card_width = Emu((CONTENT_WIDTH - METRIC_GAP * (len(metrics) - 1)) // len(metrics))
    _set_column_widths(
        table, [card_width if index % 2 == 0 else METRIC_GAP for index in range(cols)]
    )
    card_cells = table.rows[0].cells[::2]
    for cell, (label, value) in zip(card_cells, metrics):
        holder = cell.paragraphs[0]
        _tighten(holder)
        label_paragraph, value_paragraph = _rounded_box(
            holder,
            width=card_width,
            height=METRIC_CARD_HEIGHT,
            fill=_accent_top_fill(SURFACE_HEX, _theme.get().accent, METRIC_ACCENT_PCT),
            outline=BORDER_HEX,
            radius=METRIC_CARD_RADIUS,
            insets=(Mm(3.2), Mm(3.2), Mm(3.2), Mm(2)),
            paragraphs=2,
        )
        _letter_spacing(_run(label_paragraph, label.upper(), size=8, color=MUTED), 0.5)
        _run(value_paragraph, value, size=16, bold=True, color=_primary())

    _spacer(doc)
    details = [
        ("Website read on", format_display(report.created_at, settings.report_timezone)),
        ("Report generated", generated_at_label()),
        ("Shareable link", share_url(report.slug)),
    ]
    holder = doc.add_paragraph()
    _tighten(holder)
    holder.paragraph_format.space_after = Pt(14)
    rows = _rounded_box(
        holder,
        width=CONTENT_WIDTH,
        height=DETAILS_BOX_HEIGHT,
        fill=_solid_fill(SURFACE_HEX),
        outline=None,
        radius=DETAILS_BOX_RADIUS,
        insets=(Mm(3.2), Mm(2.6), Mm(3.2), Mm(2.6)),
        paragraphs=len(details),
    )
    for row, (label, value) in zip(rows, details):
        row.paragraph_format.tab_stops.add_tab_stop(DETAILS_VALUE_OFFSET)
        row.paragraph_format.space_before = Pt(3.5)
        row.paragraph_format.space_after = Pt(3.5)
        _run(row, f"{label}\t", bold=True, color=MUTED)
        _run(row, value)


def _build_summary(doc: DocxDocument, report: Report) -> None:
    totals = compute_totals(report)
    heading = doc.add_heading("Summary", level=1)
    _set_paragraph_border(heading, "bottom", _theme.get().accent, size=12)

    plural = "s" if report.page_count != 1 else ""
    lead = doc.add_paragraph()
    _run(
        lead,
        f"Overview of the {report.page_count} page{plural} read from this website. "
        "Each page is described in detail in the sections that follow.",
        color=MUTED,
    )

    headers = ("#", "Page", "Description", "Words", "Links", "Images")
    table = _new_table(doc, rows=1, cols=len(headers))
    header_row = table.rows[0]
    _repeat_as_header(header_row)
    for index, (cell, text) in enumerate(zip(header_row.cells, headers)):
        _shade(cell, _theme.get().primary)
        _run(cell.paragraphs[0], text, size=9, bold=True, color=WHITE)
        _align_numeric(cell, index)

    for row_index, page in enumerate(report.pages):
        cells = table.add_row().cells
        _keep_row_together(table.rows[-1])
        values = (
            str(page.position),
            None,
            page_description(page),
            f"{page.word_count:,}",
            f"{page.link_count:,}",
            f"{page.image_count:,}",
        )
        for index, (cell, value) in enumerate(zip(cells, values)):
            if row_index % 2 == 1:
                _shade(cell, SURFACE_HEX)
            if value is not None:
                _run(cell.paragraphs[0], value, size=9)
            _align_numeric(cell, index)
        page_cell = cells[1].paragraphs[0]
        page_cell.paragraph_format.space_after = Pt(0)
        _run(page_cell, page.title, size=9, bold=True, color=_primary())
        _run(cells[1].add_paragraph(), page.url, size=7.5, color=MUTED)

    total_cells = table.add_row().cells
    merged = total_cells[0].merge(total_cells[2])
    _run(merged.paragraphs[0], "Total", size=9, bold=True, color=_primary())
    for index, value in ((3, totals.words), (4, totals.links), (5, totals.images)):
        _run(total_cells[index].paragraphs[0], f"{value:,}", size=9, bold=True, color=_primary())
        _align_numeric(total_cells[index], index)
    for cell in table.rows[-1].cells:
        _shade(cell, SURFACE_STRONG_HEX)
        _set_cell_border(cell, "top", _theme.get().primary, size=12)

    _set_column_widths(table, SUMMARY_COLUMN_WIDTHS, skip_last_row=True)
    merged.width = Emu(sum(SUMMARY_COLUMN_WIDTHS[:3]))
    for index in (3, 4, 5):
        total_cells[index].width = SUMMARY_COLUMN_WIDTHS[index]


def _build_charts(doc: DocxDocument, report: Report) -> None:
    charts = build_charts(report, "png")
    if not charts:
        return

    heading = doc.add_heading("Charts", level=1)
    heading.paragraph_format.space_before = Pt(20)
    heading.paragraph_format.keep_with_next = True
    _set_paragraph_border(heading, "bottom", _theme.get().accent, size=12)
    lead = doc.add_paragraph()
    lead.paragraph_format.keep_with_next = True
    _run(lead, "A visual comparison of the pages read from this website.", color=MUTED)

    for chart in charts:
        title = doc.add_paragraph()
        title.paragraph_format.space_after = Pt(1)
        title.paragraph_format.keep_with_next = True
        _run(title, chart.title, size=11, bold=True, color=_primary())
        _add_chart_image(doc, chart, CONTENT_WIDTH)
        doc.add_paragraph().paragraph_format.space_after = Pt(8)


def _add_chart_image(doc: DocxDocument, chart: ReportChart, width: int) -> None:
    caption = doc.add_paragraph()
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True
    _run(caption, chart.caption, size=9, color=MUTED)
    picture = doc.add_paragraph()
    _tighten(picture)
    picture.add_run().add_picture(BytesIO(chart.image), width=width)


def _build_page_section(doc: DocxDocument, page: ReportPage) -> None:
    settings = get_settings()
    # Word ignores page breaks inside tables and merges adjacent tables, so a collapsed
    # paragraph carries the break and keeps this table apart from the previous one.
    page_break = doc.add_paragraph()
    page_break.paragraph_format.page_break_before = True
    _collapse_paragraph(page_break)
    page_break.paragraph_format.space_after = PAGE_TITLE_TOP_SPACE

    heading_table = _new_table(doc, rows=1, cols=2, borders=False)
    _set_table_borders(heading_table, None, margins=(0, 140, 0, 0))
    _set_column_widths(
        heading_table, (PAGE_BADGE_WIDTH, Emu(CONTENT_WIDTH - PAGE_BADGE_WIDTH))
    )
    badge_cell, title_cell = heading_table.rows[0].cells
    for cell in (badge_cell, title_cell):
        _set_cell_border(cell, "bottom", _theme.get().accent, size=12)

    holder = badge_cell.paragraphs[0]
    _tighten(holder)
    (badge,) = _rounded_box(
        holder,
        width=PAGE_BADGE_SIZE[0],
        height=PAGE_BADGE_SIZE[1],
        fill=_solid_fill(_theme.get().primary),
        outline=None,
        radius=PAGE_BADGE_RADIUS,
        insets=(0, 0, 0, 0),
        anchor="ctr",
    )
    badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(badge, f"{page.position:02d}", size=14, bold=True, color=WHITE)

    title = title_cell.paragraphs[0]
    title.style = "Heading 1"
    title.paragraph_format.space_after = Pt(2)
    _run(title, page.title, size=15, bold=True, color=_primary())
    url = title_cell.add_paragraph()
    url.paragraph_format.space_after = Pt(0)
    _run(url, page.url, size=9, color=MUTED)
    _spacer(doc)

    stats = (
        ("Words", f"{page.word_count:,}"),
        ("Links", f"{page.link_count:,}"),
        ("Images", f"{page.image_count:,}"),
        ("Read on", format_display(page.fetched_at, settings.report_timezone)),
    )
    table = _new_table(doc, rows=2, cols=len(stats))
    _set_column_widths(table, [Emu(int(CONTENT_WIDTH * share)) for share in STATS_COLUMN_SHARES])
    for cell, (label, _) in zip(table.rows[0].cells, stats):
        _shade(cell, SURFACE_STRONG_HEX)
        _letter_spacing(_run(cell.paragraphs[0], label.upper(), size=8, bold=True, color=MUTED), 0.5)
    for cell, (_, value) in zip(table.rows[1].cells, stats):
        _run(cell.paragraphs[0], value, size=11, bold=True, color=_primary())

    breakdown = build_page_chart(page, "png", _theme.get())
    if breakdown:
        doc.add_heading(breakdown.title, level=2).paragraph_format.keep_with_next = True
        _add_chart_image(doc, breakdown, CONTENT_WIDTH)

    doc.add_heading("Meta description", level=2)
    if page.meta_description:
        doc.add_paragraph(page.meta_description)
    else:
        _run(doc.add_paragraph(), "This page does not provide a meta description.", italic=True, color=MUTED)

    doc.add_heading("Key headings", level=2)
    if page.headings:
        for item in page.headings:
            level = int(item.get("level", 1))
            line = doc.add_paragraph()
            fmt = line.paragraph_format
            fmt.space_before = Pt(4)
            fmt.space_after = Pt(4)
            # H2 is indented with a tab, not paragraph indent, so every row shares the same
            # indent and Word draws the dashed separators full width between all of them.
            fmt.tab_stops.add_tab_stop(HEADING_H2_INDENT)
            _set_paragraph_border(line, "bottom", BORDER_HEX, size=6, space=3, style="dashed")
            _set_paragraph_border(line, "between", BORDER_HEX, size=6, space=3, style="dashed")
            if level > 1:
                line.add_run("\t")
            _tag(line, f"H{level}", _theme.get().accent if level == 1 else MUTED_HEX)
            _run(line, "  " + item.get("text", ""))
    else:
        _run(doc.add_paragraph(), "No H1 or H2 headings were found on this page.", italic=True, color=MUTED)

    doc.add_heading("Content summary", level=2)
    # A single-cell table (not a text box) so a long summary can still flow onto the next page.
    box = _new_table(doc, rows=1, cols=1, borders=False)
    _set_table_borders(box, None, margins=SUMMARY_PADDING)
    cell = box.rows[0].cells[0]
    _shade(cell, SURFACE_HEX)
    _set_cell_border(cell, "left", _theme.get().accent, size=18)
    summary = cell.paragraphs[0]
    summary.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    summary.paragraph_format.space_after = Pt(0)
    summary.paragraph_format.line_spacing = 1.2
    _run(summary, page.summary)


# ---------- Low-level helpers ----------
def _primary() -> RGBColor:
    return RGBColor.from_string(_theme.get().primary)


def _accent() -> RGBColor:
    return RGBColor.from_string(_theme.get().accent)


def _run(paragraph: Paragraph, text: str, size: float | None = None, bold: bool = False,
         italic: bool = False, color: RGBColor | None = None):
    run = paragraph.add_run(text)
    if size:
        run.font.size = Pt(_theme.get().pt(size))
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def _spacer(doc: DocxDocument) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def _new_table(doc: DocxDocument, rows: int, cols: int, borders: bool = True) -> Table:
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table, BORDER_HEX if borders else None)
    _set_column_widths(table, [Emu(CONTENT_WIDTH // cols)] * cols)
    return table


def _set_table_borders(table: Table, color: str | None,
                       margins: tuple[int, int, int, int] = (70, 70, 110, 110)) -> None:
    """Set table borders (None = no borders) and cell margins (top, bottom, left, right in dxa)."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        if color and edge not in ("left", "right", "insideV"):
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:color"), color)
        else:
            element.set(qn("w:val"), "nil")
        borders.append(element)
    for existing in tbl_pr.findall(qn("w:tblBorders")) + tbl_pr.findall(qn("w:tblCellMar")):
        tbl_pr.remove(existing)
    tbl_pr.insert_element_before(borders, *_TBLPR_AFTER_BORDERS)

    cell_margins = OxmlElement("w:tblCellMar")
    for edge, value in zip(("top", "left", "bottom", "right"),
                           (margins[0], margins[2], margins[1], margins[3])):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        cell_margins.append(element)
    tbl_pr.insert_element_before(cell_margins, *_TBLPR_AFTER_CELL_MARGINS)


def _set_cell_border(cell: _Cell, edge: str, color: str, size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.insert_element_before(borders, *_TCPR_AFTER_BORDERS)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:color"), color)
    borders.append(element)


def _shading(fill: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    return shading


def _shade(cell: _Cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:shd"))
    if existing is not None:
        tc_pr.remove(existing)
    tc_pr.insert_element_before(_shading(fill), *_TCPR_AFTER_SHD)


def _tag(paragraph: Paragraph, text: str, fill: str) -> None:
    """Small rounded white-on-colour label, lowered so it sits centred on the text line."""
    (label,) = _rounded_box(
        paragraph,
        width=TAG_SIZE[0],
        height=TAG_SIZE[1],
        fill=_solid_fill(fill),
        outline=None,
        radius=TAG_RADIUS,
        insets=(0, 0, 0, 0),
        anchor="ctr",
    )
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(label, text, size=7.5, bold=True, color=WHITE)
    position = OxmlElement("w:position")
    position.set(qn("w:val"), str(TAG_BASELINE_SHIFT_HALF_PT))
    paragraph.runs[-1]._r.get_or_add_rPr().insert_element_before(position, *_RPR_AFTER_POSITION)


def _letter_spacing(run, points: float):
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(round(points * 20)))
    run._r.get_or_add_rPr().insert_element_before(spacing, *_RPR_AFTER_SPACING)
    return run


def _tighten(paragraph: Paragraph) -> None:
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0


def _solid_fill(color: str) -> str:
    return f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'


def _accent_top_fill(color: str, accent: str, accent_pct: int) -> str:
    """Hard-stop gradient: a thin accent band along the top that follows the rounded corners."""
    stops = ((0, accent), (accent_pct, accent), (accent_pct + 1, color), (100_000, color))
    gs = "".join(f'<a:gs pos="{pos}"><a:srgbClr val="{value}"/></a:gs>' for pos, value in stops)
    return f'<a:gradFill rotWithShape="1"><a:gsLst>{gs}</a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>'


def _rounded_box(holder: Paragraph, width: int, height: int, fill: str, outline: str | None,
                 radius: int, insets: tuple[int, int, int, int], anchor: str = "t",
                 paragraphs: int = 1) -> list[Paragraph]:
    """Add an inline rounded-rectangle text box to `holder`; returns its (editable) paragraphs."""
    shape_id = 1000 + len(holder.part.element.findall(".//" + qn("wp:docPr")))
    line = (
        f'<a:ln w="9525"><a:solidFill><a:srgbClr val="{outline}"/></a:solidFill></a:ln>'
        if outline else "<a:ln><a:noFill/></a:ln>"
    )
    left, top, right, bottom = (int(value) for value in insets)
    xml = (
        f'<w:drawing xmlns:w="{nsmap["w"]}" xmlns:wp="{nsmap["wp"]}" xmlns:a="{nsmap["a"]}" '
        f'xmlns:wps="{WPS_NS}">'
        '<wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="{int(width)}" cy="{int(height)}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:docPr id="{shape_id}" name="Box {shape_id}"/>'
        "<wp:cNvGraphicFramePr/>"
        f'<a:graphic><a:graphicData uri="{WPS_NS}"><wps:wsp><wps:cNvSpPr/>'
        f'<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{int(width)}" cy="{int(height)}"/></a:xfrm>'
        f'<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val {radius}"/></a:avLst></a:prstGeom>'
        f"{fill}{line}</wps:spPr>"
        "<wps:txbx><w:txbxContent>" + "<w:p/>" * paragraphs + "</w:txbxContent></wps:txbx>"
        f'<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="{left}" tIns="{top}" rIns="{right}" '
        f'bIns="{bottom}" anchor="{anchor}" anchorCtr="0"><a:noAutofit/></wps:bodyPr>'
        "</wps:wsp></a:graphicData></a:graphic></wp:inline></w:drawing>"
    )
    drawing = parse_xml(xml)
    holder.add_run()._r.append(drawing)
    content = drawing.find(".//" + qn("w:txbxContent"))
    boxes = [Paragraph(p, holder._parent) for p in content.findall(qn("w:p"))]
    for box in boxes:
        _tighten(box)
    return boxes


def _collapse_paragraph(paragraph: Paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(1)


def _set_column_widths(table: Table, widths, skip_last_row: bool = False) -> None:
    """Word lays out by the table grid (gridCol), so set it as well as every cell's width."""
    for column, width in zip(table.columns, widths):
        column.width = width
    rows = table.rows[:-1] if skip_last_row else table.rows
    for row in rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def _align_numeric(cell: _Cell, column_index: int) -> None:
    if column_index == 0:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif column_index >= 3:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _repeat_as_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def _keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:cantSplit")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def _set_paragraph_border(paragraph: Paragraph, edge: str, color: str = BORDER_HEX,
                          size: int = 6, space: int = 4, style: str = "single") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.insert_element_before(borders, *_PPR_AFTER_PBDR)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), style)
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), color)
    borders.append(element)


def _add_field(paragraph: Paragraph, instruction: str) -> None:
    """Insert a Word field (e.g. PAGE, NUMPAGES) that Word evaluates when the document opens."""
    def fld_char(kind: str):
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        return element

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction

    for part in (fld_char("begin"), instr, fld_char("separate")):
        run = _run(paragraph, "", size=8, color=MUTED)
        run._r.append(part)
    _run(paragraph, "1", size=8, color=MUTED)
    _run(paragraph, "", size=8, color=MUTED)._r.append(fld_char("end"))
