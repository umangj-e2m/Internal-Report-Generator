from io import BytesIO

from docx import Document
from pptx import Presentation

from services.shared.constants import DOCX_MEDIA_TYPE, PDF_MEDIA_TYPE, PPTX_MEDIA_TYPE
from tests.fixtures.sample_site import BASE


def test_generate_then_view_and_download_every_format(client):
    report = client.post("/api/reports", json={"url": f"{BASE}/"}).json()
    slug = report["slug"]

    html = client.get(f"/api/reports/{slug}/html")
    assert html.status_code == 200
    assert html.headers["content-type"].startswith("text/html")
    assert "Welcome to Example Co" in html.text

    docx = client.get(f"/api/reports/{slug}/docx")
    assert docx.status_code == 200
    assert docx.headers["content-type"] == DOCX_MEDIA_TYPE
    assert docx.headers["content-disposition"] == f'attachment; filename="{slug}.docx"'
    assert len(Document(BytesIO(docx.content)).tables) >= 3

    slides = client.get(f"/api/reports/{slug}/slides")
    assert slides.status_code == 200
    assert 'class="slide slide--title"' in slides.text

    pptx = client.get(f"/api/reports/{slug}/pptx")
    assert pptx.status_code == 200
    assert pptx.headers["content-type"] == PPTX_MEDIA_TYPE
    assert pptx.headers["content-disposition"] == f'attachment; filename="{slug}.pptx"'
    assert len(Presentation(BytesIO(pptx.content)).slides) >= 5

    pdf_inline = client.get(f"/api/reports/{slug}/pdf")
    assert pdf_inline.status_code == 200
    assert pdf_inline.headers["content-type"] == PDF_MEDIA_TYPE
    assert pdf_inline.headers["content-disposition"] == f'inline; filename="{slug}.pdf"'
    assert pdf_inline.content.startswith(b"%PDF")

    pdf_download = client.get(f"/api/reports/{slug}/pdf", params={"download": "true"})
    assert pdf_download.headers["content-disposition"] == f'attachment; filename="{slug}.pdf"'
