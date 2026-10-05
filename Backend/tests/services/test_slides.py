from io import BytesIO

from pptx import Presentation

from services.exports.html_export import render_slides_html
from services.exports.pptx_export import render_report_pptx
from services.exports.slides import BODY_HEIGHT, META_ON_OVERVIEW_CHARS, build_deck


def _page_content(deck, position: int) -> list[dict]:
    return [
        section
        for slide in deck.slides
        if slide.kind == "page_content" and slide.data["position"] == position
        for section in slide.data["sections"]
    ]


def test_deck_starts_with_title_ends_with_closing_and_is_numbered(sample_report):
    deck = build_deck(sample_report, "png")
    kinds = [slide.kind for slide in deck.slides]

    assert kinds[:3] == ["title", "overview", "summary"]
    assert kinds[-1] == "closing"
    assert kinds.count("page") == sample_report.page_count
    assert [slide.number for slide in deck.slides] == list(range(1, len(deck.slides) + 1))


def test_many_headings_continue_on_extra_slides(sample_report):
    page = sample_report.pages[0]
    page.headings = [{"level": 2, "text": f"Heading number {index}"} for index in range(40)]

    deck = build_deck(sample_report, "png")
    sections = [s for s in _page_content(deck, page.position) if s["kind"] == "headings"]

    assert len(sections) > 1
    assert sections[1]["title"] == "Key headings (continued)"
    assert [item["text"] for s in sections for item in s["items"]] == [
        heading["text"] for heading in page.headings
    ]


def test_long_summary_is_split_without_losing_words(sample_report):
    page = sample_report.pages[0]
    page.summary = " ".join(f"word{index}" for index in range(1500))

    deck = build_deck(sample_report, "png")
    sections = [s for s in _page_content(deck, page.position) if s["title"].startswith("Content")]

    assert len(sections) > 1
    assert " ".join(section["text"] for section in sections) == page.summary
    for slide in deck.slides:
        if slide.kind == "page_content":
            assert all(s["body_top"] + s["body_height"] <= BODY_HEIGHT + 1e-9
                       for s in slide.data["sections"])


def test_larger_text_needs_more_content_slides(sample_report):
    page = sample_report.pages[0]
    page.summary = " ".join(f"word{index}" for index in range(1500))

    def summary_parts(size: str) -> list[dict]:
        sample_report.font_size = size
        deck = build_deck(sample_report, "png")
        return [s for s in _page_content(deck, page.position) if s["title"].startswith("Content")]

    small, large = summary_parts("small"), summary_parts("large")

    assert len(large) > len(small)
    assert " ".join(section["text"] for section in large) == page.summary


def test_long_meta_description_moves_to_content_slides(sample_report):
    page = sample_report.pages[0]
    page.meta_description = "Long description " * (META_ON_OVERVIEW_CHARS // 10)

    deck = build_deck(sample_report, "png")
    overview = next(s for s in deck.slides if s.kind == "page" and s.data["position"] == 1)

    assert overview.data["meta"] is None
    assert _page_content(deck, 1)[0]["title"] == "Meta description"


def test_pptx_and_html_have_the_same_slide_count(sample_report):
    presentation = Presentation(BytesIO(render_report_pptx(sample_report)))
    html = render_slides_html(sample_report)

    assert len(presentation.slides) == html.count('<section class="slide ')
    assert "Sample &lt;Site&gt;" in html
